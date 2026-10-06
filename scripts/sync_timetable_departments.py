"""Join public VKU timetable to supplied teacher names without reading passwords."""
import argparse
from collections import Counter, defaultdict
import hashlib
from html.parser import HTMLParser
import json
from pathlib import Path
import re
import sqlite3
import unicodedata

import openpyxl
from backend.db import Database, audit, now
from scripts.provision_teachers import plain_name


def name_key(name, accents=True):
    value = ' '.join(plain_name(unicodedata.normalize('NFC', name)).casefold().split())
    if not accents:
        value = ''.join(c for c in unicodedata.normalize('NFD', value.replace('đ', 'd'))
                        if not unicodedata.combining(c))
    return value


class TimetableParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.active = False
        self.cell = None
        self.rows = []
        self.cells = []
        self.select = None
        self.option = None
        self.period = {}

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'select':
            self.select = next((c for c in ('hocky', 'namhoc') if c in attrs.get('class', '').split()), None)
        if tag == 'option' and self.select and 'selected' in attrs:
            self.option = [self.select, attrs.get('value'), '']
        if tag == 'tbody' and 'lophocphan' in attrs.get('class', '').split():
            self.active = True
        if self.active and tag == 'tr':
            self.cells = []
        if self.active and tag == 'td':
            self.cell = {'text': '', 'rowspan': int(attrs.get('rowspan', '1')), 'href': ''}
        if self.cell is not None and tag == 'a':
            self.cell['href'] = attrs.get('href', '')
        if self.cell is not None and tag == 'br':
            self.cell['text'] += ' '

    def handle_data(self, data):
        if self.cell is not None:
            self.cell['text'] += data
        if self.option is not None:
            self.option[2] += data

    def handle_endtag(self, tag):
        if tag == 'option' and self.option is not None:
            key, value, label = self.option
            self.period[key] = (value, ' '.join(label.split()))
            self.option = None
        if tag == 'select':
            self.select = None
        if tag == 'td' and self.cell is not None:
            self.cell['text'] = ' '.join(self.cell['text'].split())
            self.cells.append(self.cell)
            self.cell = None
        if tag == 'tr' and self.active and self.cells:
            self.rows.append(self.cells)
        if tag == 'tbody':
            self.active = False


def read_timetable(path, year, semester):
    parser = TimetableParser()
    parser.feed(Path(path).read_text(encoding='utf-8'))
    match = re.search(r'(\d{4})\s*-\s*(\d{4})', parser.period.get('namhoc', ('', ''))[1])
    if not match or '-'.join(match.groups()) != year or parser.period.get('hocky', ('', ''))[0] != semester:
        raise ValueError('Học kỳ/năm của trang nguồn không khớp kỳ yêu cầu.')
    carry, rows = {}, []
    for cells in parser.rows:
        expanded = []
        incoming = iter(cells)
        for col in range(8):
            if col in carry:
                cell, remaining = carry[col]
                if remaining == 1:
                    del carry[col]
                else:
                    carry[col] = (cell, remaining - 1)
            else:
                cell = next(incoming, None)
                if cell is None:
                    raise ValueError('Cấu trúc cột thời khóa biểu thay đổi.')
                if cell['rowspan'] > 1:
                    carry[col] = (cell, cell['rowspan'] - 1)
            expanded.append(cell)
        if next(incoming, None) is not None:
            raise ValueError('Thừa cột thời khóa biểu.')
        text = [c['text'] for c in expanded]
        rows.append(dict(zip(('stt', 'course', 'teacher', 'schedule', 'room', 'weeks', 'students'), text[:7]),
                         source_id=expanded[7]['href'].rsplit('/', 1)[-1]))
    if not rows or carry:
        raise ValueError('Thời khóa biểu rỗng hoặc rowspan chưa hoàn tất.')
    return rows


def read_teacher_accounts(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        units = {str(r[1]).strip(): str(r[2]).strip() for r in workbook['Danh muc Khoa'].iter_rows(values_only=True)
                 if r[0] != 'STT' and r[1] and r[2]}
        sheet = workbook['Tai khoan VKU']
        records = {}
        started = False
        for row in sheet.iter_rows(values_only=True):
            if row[2] == 'Email đăng nhập':
                started = True
                continue
            if not started or not row[2] or 'GV' not in [r.strip() for r in str(row[4]).split(',')]:
                continue
            # Ignore the password column entirely, including formulas there.
            if any(str(row[i] or '').startswith('=') for i in (1, 2, 4, 5)):
                raise ValueError('Không chấp nhận công thức trong tên/email/vai trò/khoa.')
            record = dict(name=str(row[1]).strip(), email=str(row[2]).strip().lower(), department=str(row[5]).strip())
            if record['email'] in records and records[record['email']] != record:
                raise ValueError('Một email có thông tin GV mâu thuẫn.')
            records[record['email']] = record
        if not records:
            raise ValueError('Excel không có tài khoản GV.')
        return list(records.values()), units
    finally:
        workbook.close()


def plan_import(database, rows, accounts, units, year, semester):
    indexes = [defaultdict(list), defaultdict(list)]
    for account in accounts:
        for i, accents in enumerate((True, False)):
            indexes[i][name_key(account['name'], accents)].append(account)
    with database.connect() as db:
        users = {u['email']: dict(u) for u in db.execute("SELECT id,email,name,department FROM users WHERE role='teacher' AND active=1")}
    grouped, unresolved = {}, []
    for row in rows:
        candidates = indexes[0].get(name_key(row['teacher']), []) or indexes[1].get(name_key(row['teacher'], False), [])
        reason = None
        if len(candidates) != 1:
            reason = 'Tên trùng' if candidates else 'Không có GV trong Excel'
        else:
            account = candidates[0]
            user = users.get(account['email'])
            if account['department'] not in units:
                reason = 'Khoa chưa xác định trong Excel'
            elif not user or name_key(user['name'], False) != name_key(account['name'], False):
                reason = 'Tài khoản GV chưa hoạt động hoặc họ tên không khớp'
        if reason:
            unresolved.append({'teacher': row['teacher'], 'course': row['course'], 'reason': reason})
            continue
        key = (account['email'], row['course'])
        if key not in grouped:
            code = 'TKB-' + hashlib.sha256(('|'.join((account['email'], year, semester, row['course']))).encode()).hexdigest()[:10].upper()
            grouped[key] = dict(code=code, title=row['course'], year=year, semester=semester,
                                department=account['department'], teacher_id=user['id'], email=account['email'], assignments=[], source_ids=[])
        assignment = (row['schedule'], row['room'], row['weeks'])
        if assignment not in grouped[key]['assignments']:
            grouped[key]['assignments'].append(assignment)
        if row.get('source_id') and row['source_id'] not in grouped[key]['source_ids']:
            grouped[key]['source_ids'].append(row['source_id'])
    # Identical titles assigned to multiple people can be shared classes or separate
    # projects without a group number. Report them for review; never silently merge.
    titles = defaultdict(set)
    for course in grouped.values():
        titles[course['title']].add(course['email'])
    shared = [{'title': title, 'teachers': sorted(emails)} for title, emails in titles.items() if len(emails) > 1]
    return {'courses': list(grouped.values()), 'unresolved': unresolved, 'shared_titles': shared,
            'counts': dict(Counter(c['department'] for c in grouped.values())), 'source_rows': len(rows)}


def apply_plan(database, plan):
    created = updated = 0
    with database.connect() as db:
        db.execute('BEGIN IMMEDIATE')
        for course in plan['courses']:
            teacher = db.execute("SELECT * FROM users WHERE id=? AND email=? AND role='teacher' AND active=1",
                                 (course['teacher_id'], course['email'])).fetchone()
            if not teacher:
                raise ValueError('Quyền GV đã thay đổi sau đối chiếu; cần chạy lại.')
            existing = db.execute('SELECT * FROM courses WHERE code=? AND year=? AND semester=?',
                                  (course['code'], course['year'], course['semester'])).fetchone()
            schedule, room, weeks = ('\n'.join(a[i] for a in course['assignments']) for i in range(3))
            values = (course['department'], course['teacher_id'], schedule, room, weeks)
            if existing:
                if existing['teacher_id'] != course['teacher_id'] or existing['title'] != course['title']:
                    raise ValueError('Mã lớp đã thuộc phân công khác; không ghi đè.')
                if tuple(existing[k] for k in ('department', 'teacher_id', 'teaching_schedule', 'teaching_room', 'teaching_weeks')) != values:
                    db.execute('UPDATE courses SET department=?,teacher_id=?,teaching_schedule=?,teaching_room=?,teaching_weeks=? WHERE id=?', values + (existing['id'],))
                    updated += 1
            else:
                db.execute('INSERT INTO courses(code,title,year,semester,department,teacher_id,teaching_schedule,teaching_room,teaching_weeks) VALUES(?,?,?,?,?,?,?,?,?)',
                           (course['code'], course['title'], course['year'], course['semester']) + values)
                created += 1
        audit(db, None, 'sync_timetable_departments', 'courses', json.dumps({'created': created, 'updated': updated, 'counts': plan['counts']}))
    return {'created': created, 'updated': updated}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--html', required=True)
    parser.add_argument('--accounts', required=True)
    parser.add_argument('--year', required=True)
    parser.add_argument('--semester', choices=('1', '2', '3'), required=True)
    parser.add_argument('--database', default='runtime/gradebook.sqlite3')
    parser.add_argument('--report', required=True)
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    Path(args.report).parent.mkdir(parents=True, exist_ok=True)
    database = Database(args.database)
    rows = read_timetable(args.html, args.year, args.semester)
    accounts, units = read_teacher_accounts(args.accounts)
    plan = plan_import(database, rows, accounts, units, args.year, args.semester)
    if args.apply:
        backup = Path(args.report).with_suffix('.backup.sqlite3')
        if backup.exists():
            raise ValueError('Bản sao lưu đã tồn tại; chọn report mới để không ghi đè.')
        with database.connect() as source, sqlite3.connect(backup) as target:
            source.backup(target)
        plan['applied'] = apply_plan(database, plan)
    plan.update(year=args.year, semester=args.semester, source_url='https://daotao.vku.udn.vn/thoi-khoa-bieu', units=units,
                imported_at=now(), html_sha256=hashlib.sha256(Path(args.html).read_bytes()).hexdigest(),
                accounts_sha256=hashlib.sha256(Path(args.accounts).read_bytes()).hexdigest())
    Path(args.report).write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: plan[k] for k in ('counts', 'source_rows')}, ensure_ascii=False))
    print('Unresolved rows:', len(plan['unresolved']), '| Shared titles:', len(plan['shared_titles']), '|', plan.get('applied', 'preview'))


if __name__ == '__main__':
    main()
